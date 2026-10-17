"""Loaders for PDF (page numbers), DOCX, TXT, MD (AI-RG-01/03). Every failure becomes an IngestionError with a clear reason."""
from __future__ import annotations

import logging
import time
from dataclasses import dataclass, field
from io import BytesIO

from app.rag.errors import IngestionError

log = logging.getLogger(__name__)


@dataclass
class PageText:
    page: int | None   # 1-based for PDF, None when the format has no pages
    text: str


@dataclass
class LoadedDocument:
    pages: list[PageText]
    page_count: int | None
    warnings: list[str] = field(default_factory=list)

    @property
    def total_chars(self) -> int:
        return sum(len(p.text) for p in self.pages)


def _check_deadline(deadline: float | None) -> None:
    if deadline is not None and time.monotonic() > deadline:
        raise IngestionError("Ingestion took too long and was stopped. Try a smaller document.")


def load_document(
    data: bytes, source_type: str, *, max_pages: int = 500, max_chars: int = 3_000_000, deadline: float | None = None
) -> LoadedDocument:
    st = source_type.upper()
    if not data:
        raise IngestionError("The file is empty.")
    if st == "PDF":
        doc = _load_pdf(data, max_pages, max_chars, deadline)
    elif st == "DOCX":
        doc = _load_docx(data, max_chars)
    elif st in ("TXT", "MD"):
        doc = _load_text(data)
    else:
        raise IngestionError(f"Unsupported source type: {source_type}")
    if doc.total_chars > max_chars:
        raise IngestionError(f"The document has too much text (limit {max_chars:,} characters).")
    if not any(p.text.strip() for p in doc.pages):
        raise IngestionError("The document contains no readable text.")
    return doc


# ----------------------------------------------------------------------------- PDF
def _load_pdf(data: bytes, max_pages: int, max_chars: int, deadline: float | None) -> LoadedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    if not data.lstrip()[:1024].startswith(b"%PDF"):
        raise IngestionError("The file is not a valid PDF (missing PDF header).")
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")  # many "encrypted" PDFs only have an owner password
            except Exception:  # noqa: BLE001 - e.g. missing crypto backend
                ok = 0
            if not ok:
                raise IngestionError("The PDF is password-protected. Remove the password and upload it again.")
        n = len(reader.pages)
    except IngestionError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError, RecursionError, TypeError, AttributeError) as exc:
        log.info("pdf unreadable: %s", type(exc).__name__)
        raise IngestionError("The PDF is corrupt or unreadable.") from exc
    except Exception as exc:  # noqa: BLE001
        log.info("pdf unreadable: %s", type(exc).__name__)
        raise IngestionError("The PDF is corrupt or unreadable.") from exc

    if n == 0:
        raise IngestionError("The PDF has no pages.")
    if n > max_pages:
        raise IngestionError(f"The PDF has {n} pages; the limit is {max_pages}.")

    pages: list[PageText] = []
    warnings: list[str] = []
    failed = chars = 0
    for i in range(n):
        _check_deadline(deadline)
        try:
            text = reader.pages[i].extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            failed += 1
            text = ""
            log.info("pdf page %d unreadable: %s", i + 1, type(exc).__name__)
        chars += len(text)
        if chars > max_chars:
            raise IngestionError(f"The document has too much text (limit {max_chars:,} characters).")
        pages.append(PageText(i + 1, text))
    if failed:
        warnings.append(f"{failed} of {n} pages could not be read and were skipped.")
    blank = sum(1 for p in pages if len(p.text.strip()) < 15)
    if chars < 20 * max(1, n) and blank == n:
        raise IngestionError(
            "No extractable text: this looks like a scanned or image-only PDF (OCR is not enabled)."
        )
    if blank > n // 2:
        warnings.append(f"{blank} of {n} pages have little or no text (scanned pages are not OCR'd).")
    return LoadedDocument(pages, n, warnings)


# ----------------------------------------------------------------------------- DOCX
def _load_docx(data: bytes, max_chars: int) -> LoadedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    if not data.startswith(b"PK"):
        raise IngestionError("The file is not a valid DOCX.")
    try:
        doc = Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        log.info("docx unreadable: %s", type(exc).__name__)
        raise IngestionError("The DOCX is corrupt or unreadable.") from exc

    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                lines.append("")
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower() == "title":
                lines.append(f"# {text}")
            elif style.lower().startswith("heading"):
                digits = "".join(c for c in style if c.isdigit())
                level = min(max(int(digits) if digits else 1, 1), 6)
                lines.append(f"{'#' * level} {text}")
            else:
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells = []
                for c in row.cells:
                    t = " ".join(c.text.split())
                    if not cells or t != cells[-1]:   # merged cells repeat their text
                        cells.append(t)
                if any(cells):
                    lines.append(" | ".join(cells))
            lines.append("")
        if sum(len(x) for x in lines) > max_chars * 2:
            raise IngestionError(f"The document has too much text (limit {max_chars:,} characters).")
    return LoadedDocument([PageText(None, "\n".join(lines))], None, [])


# ----------------------------------------------------------------------------- TXT / MD
def _load_text(data: bytes) -> LoadedDocument:
    warnings: list[str] = []
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
        warnings.append("The file is not valid UTF-8; it was read as Latin-1.")
    if "\x00" in text[:4096]:
        raise IngestionError("The file looks binary, not text.")
    return LoadedDocument([PageText(None, text)], None, warnings)
