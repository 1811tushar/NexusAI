"""Tiny builders for test documents (no binary fixtures in git)."""
from __future__ import annotations

from io import BytesIO


def make_pdf(pages: list[str]) -> bytes:
    """Minimal valid multi-page text PDF (Helvetica). Empty strings give pages with no text (like a scan)."""
    objs: dict[int, bytes] = {}
    n = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objs[2] = f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode()
    objs[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for i, text in enumerate(pages):
        page_no, content_no = 4 + 2 * i, 5 + 2 * i
        ops = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
        for line in text.split("\n") if text else []:
            esc = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({esc}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1", "replace")
        objs[page_no] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_no} 0 R "
                         f"/Resources << /Font << /F1 3 0 R >> >> >>").encode()
        objs[content_no] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    out = BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objs):
        offsets[num] = out.tell()
        out.write(b"%d 0 obj\n" % num + objs[num] + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for num in sorted(objs):
        out.write(b"%010d 00000 n \n" % offsets[num])
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def encrypt_pdf(pdf: bytes, password: str = "secret") -> bytes:
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter(clone_from=PdfReader(BytesIO(pdf)))
    w.encrypt(password, algorithm="RC4-128")
    out = BytesIO()
    w.write(out)
    return out.getvalue()


def make_docx(blocks: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    """blocks: (style, text) with style in {'Title','Heading 1','Heading 2','Normal'}."""
    from docx import Document

    d = Document()
    for style, text in blocks:
        if style == "Title":
            d.add_heading(text, level=0)
        elif style.startswith("Heading"):
            d.add_heading(text, level=int(style.split()[-1]))
        else:
            d.add_paragraph(text)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    out = BytesIO()
    d.save(out)
    return out.getvalue()
